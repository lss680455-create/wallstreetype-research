#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md_to_docx.py — Wall Street-style Markdown → Word (.docx) converter.

Renders an equity research report written against templates/report_template.md
into a Word document that follows sell-side conventions (Goldman Sachs /
Morgan Stanley / J.P. Morgan style): cover page with rating & target-price
box, Executive Summary, numbered sections, compact financial tables,
Exhibit-numbered charts, analyst block, conflict-of-interest line, risk &
disclosure pages, and a "Page X of Y" footer.

  Usage:
    python md_to_docx.py report.md                    # report.docx next to it
    python md_to_docx.py report.md -o out/report.docx
    python md_to_docx.py report.md --pdf              # also export PDF (Word)
    python md_to_docx.py report.md --toc              # insert a TOC field
    python md_to_docx.py report.md --no-numbering     # sections unnumbered
    python md_to_docx.py report.md --style goldman_hardline   # institution theme
    python md_to_docx.py report.md --style my_theme.json      # custom theme file

  Dependencies (pip):
    python-docx            (required)
    markdown               (required, for md -> HTML)
    docx2pdf               (optional, only for --pdf; needs Microsoft Word)

  Design notes
  ------------
  * The Markdown file is the MASTER. docx is for editing/submission, PDF for
    distribution. GitHub renders the .md; the docx/PDF are build artifacts.
  * Frontmatter (YAML between the first two `---` lines) drives the cover,
    rating box, analyst block, header/footer and disclosure line.
  * Body conventions (see report_template.md): H1 = numbered major section,
    `![Caption](path)` = Exhibit-numbered chart, a line starting with
    "Source:" after a table/chart = small gray source note, a line directly
    above a table matching "Exhibit N:" / "Table N:" / "Chart N:" = caption.
  * English-first, Chinese-capable: body font Calibri with 等线 (DengXian)
    registered as the East-Asian face.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import markdown as md_lib
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ----------------------------------------------------------------------------
# Palette (Wall Street: navy + gray, restrained accent)
# ----------------------------------------------------------------------------
NAVY      = RGBColor(0x0B, 0x25, 0x45)   # headings, company name
ACCENT    = RGBColor(0x1F, 0x38, 0x64)   # sub-headings
DARKGRAY  = RGBColor(0x33, 0x33, 0x33)
MIDGRAY   = RGBColor(0x59, 0x59, 0x59)
LIGHTGRAY = RGBColor(0x8A, 0x8A, 0x8A)
RULEGRAY  = "BFBFBF"
SHADE_HEAD= "1F3864"                     # table header fill (navy)
SHADE_ALT = "F2F4F8"                     # zebra / rating-box fill
SHADE_CODE= "F5F5F5"
EASTASIA  = "等线"                        # DengXian, for Chinese glyphs

RATING_COLORS = {
    "overweight": RGBColor(0x1E, 0x7B, 0x34),
    "buy":        RGBColor(0x1E, 0x7B, 0x34),
    "equal-weight": RGBColor(0xB0, 0x7D, 0x2B),
    "hold":       RGBColor(0xB0, 0x7D, 0x2B),
    "neutral":    RGBColor(0xB0, 0x7D, 0x2B),
    "underweight": RGBColor(0xB0, 0x30, 0x30),
    "sell":       RGBColor(0xB0, 0x30, 0x30),
}

# ----------------------------------------------------------------------------
# Style themes — templates/styles/*.json drive fonts / colours / layout /
# tables / rating box / masthead.  With no --style the built-in default below
# reproduces the original hard-coded look exactly.
# ----------------------------------------------------------------------------
STYLES_DIR = Path(__file__).resolve().parent / "styles"

DEFAULT_THEME_DATA = {
    "id": "default",
    "display_name": "Default (built-in Wall Street look)",
    "fonts": {"body": "Calibri", "heading": "Calibri", "mono": "Consolas"},
    "colors": {
        "navy": "0B2545", "accent": "1F3864", "text": "333333",
        "midgray": "595959", "lightgray": "8A8A8A", "rule": "BFBFBF",
        "table_header_fill": "1F3864", "table_header_text": "FFFFFF",
        "table_zebra": "F2F4F8", "rating_box_fill": "F2F4F8",
        # extras used only by the legacy look (not part of the theme schema)
        "table_grid": "C9CFD8", "code_text": "8B1A1A",
        "code_fill": "F5F5F5", "link": "0563C1",
    },
    "layout": {"margin_lr": 0.95, "margin_tb": 0.85, "body_size": 10.0,
               "h1_size": 14, "h2_size": 12, "h3_size": 11,
               "h1_font": "body", "line_spacing": 1.08},
    "tables": {"vertical_rules": True, "zebra": True,
               "rule_above_header": False, "rule_below_header": False},
    "rating_box": {"style": "band", "fill": "F2F4F8", "border": True},
    "masthead": {"rule_under": True, "uppercase_meta": True},
    "rating_colors": {
        "overweight": "1E7B34", "buy": "1E7B34",
        "equal-weight": "B07D2B", "hold": "B07D2B", "neutral": "B07D2B",
        "underweight": "B03030", "sell": "B03030",
    },
}


def _hex6(value, fallback: str) -> str:
    """Normalise a colour value to a 6-digit uppercase hex string."""
    if value is None:
        return fallback
    s = str(value).strip().lstrip("#")
    if len(s) == 3:
        s = "".join(ch * 2 for ch in s)
    if len(s) != 6 or any(ch not in "0123456789abcdefABCDEF" for ch in s):
        return fallback
    return s.upper()


def _rgb6(hexstr: str) -> RGBColor:
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


class Theme:
    """Resolved theme: every value falls back to the built-in default."""

    def __init__(self, data: dict | None = None):
        data = data if isinstance(data, dict) else {}

        def section(name: str) -> dict:
            merged = dict(DEFAULT_THEME_DATA[name])
            val = data.get(name)
            if isinstance(val, dict):
                merged.update(val)
            return merged

        self.id = str(data.get("id") or DEFAULT_THEME_DATA["id"])
        self.display_name = str(data.get("display_name") or self.id)
        self.fonts = section("fonts")
        self.colors = section("colors")
        self.layout = section("layout")
        self.tables = section("tables")
        self.rating_box = section("rating_box")
        self.masthead = section("masthead")

        rc = data.get("rating_colors")
        if isinstance(rc, dict):
            self.rating_colors = {k: _rgb6(_hex6(v, "000000")) for k, v in rc.items()}
        elif data:
            # external theme without an explicit map: keep rating text neutral
            self.rating_colors = {}
        else:
            self.rating_colors = {k: _rgb6(v)
                                  for k, v in DEFAULT_THEME_DATA["rating_colors"].items()}

        # fonts
        self.body_font = str(self.fonts["body"])
        self.heading_font = str(self.fonts["heading"])
        self.mono_font = str(self.fonts["mono"])

        # colours (hex strings + RGBColor twins)
        c = self.colors
        self.navy_hex = _hex6(c.get("navy"), "0B2545")
        self.accent_hex = _hex6(c.get("accent"), "1F3864")
        self.text_hex = _hex6(c.get("text"), "333333")
        self.midgray_hex = _hex6(c.get("midgray"), "595959")
        self.lightgray_hex = _hex6(c.get("lightgray"), "8A8A8A")
        self.rule = _hex6(c.get("rule"), "BFBFBF")
        self.table_header_fill = _hex6(c.get("table_header_fill"), "1F3864")
        self.table_header_text_hex = _hex6(c.get("table_header_text"), "FFFFFF")
        self.table_zebra = _hex6(c.get("table_zebra"), "") or None
        rb_raw = data.get("rating_box") if isinstance(data.get("rating_box"), dict) else {}
        self.rating_fill = _hex6(rb_raw.get("fill"),
                                 _hex6(c.get("rating_box_fill"), "F2F4F8"))
        self.table_grid = _hex6(c.get("table_grid"), self.rule)
        self.code_fill = _hex6(c.get("code_fill"), self.table_zebra or "F5F5F5")
        self.code_text_hex = _hex6(c.get("code_text"), self.navy_hex)
        self.link_hex = _hex6(c.get("link"), self.accent_hex)

        self.navy = _rgb6(self.navy_hex)
        self.accent = _rgb6(self.accent_hex)
        self.text = _rgb6(self.text_hex)
        self.midgray = _rgb6(self.midgray_hex)
        self.lightgray = _rgb6(self.lightgray_hex)
        self.rule_rgb = _rgb6(self.rule)
        self.table_header_text = _rgb6(self.table_header_text_hex)
        self.code_text = _rgb6(self.code_text_hex)
        self.link = _rgb6(self.link_hex)

        # layout
        L = self.layout
        self.margin_lr = float(L.get("margin_lr") or 0.95)
        self.margin_tb = float(L.get("margin_tb") or 0.85)
        self.body_size = float(L.get("body_size") or 10.0)
        self.h1_size = float(L.get("h1_size") or 14)
        self.h2_size = float(L.get("h2_size") or 12)
        self.h3_size = float(L.get("h3_size") or 11)
        self.h4_size = max(self.h3_size - 0.5, 9.0)
        self.h1_font = ("heading"
                        if str(L.get("h1_font", "heading")).strip().lower() == "heading"
                        else "body")
        self.line_spacing = float(L.get("line_spacing") or 1.08)
        self.table_size = round(self.body_size - 1.5, 1)

        # tables
        self.vertical_rules = bool(self.tables.get("vertical_rules", True))
        self.zebra = bool(self.tables.get("zebra", True))
        self.rule_above_header = bool(self.tables.get("rule_above_header", False))
        self.rule_below_header = bool(self.tables.get("rule_below_header", False))

        # rating box
        self.rating_style = str(self.rating_box.get("style", "band")).strip().lower()
        self.rating_border = bool(self.rating_box.get("border", True))

        # masthead
        self.masthead_rule = bool(self.masthead.get("rule_under", True))
        self.uppercase_meta = bool(self.masthead.get("uppercase_meta", True))

    @property
    def h1_font_name(self) -> str:
        return self.heading_font if self.h1_font == "heading" else self.body_font


def available_styles() -> list[str]:
    if not STYLES_DIR.is_dir():
        return []
    return sorted(p.stem for p in STYLES_DIR.glob("*.json"))


def load_theme(style: str | None) -> Theme:
    """Resolve --style: None -> built-in default; id -> styles/<id>.json; path -> file."""
    if not style:
        return Theme({})
    raw = str(style).strip()
    cand = Path(raw)
    if not (cand.is_file() and cand.suffix.lower() == ".json"):
        for probe in (STYLES_DIR / f"{raw}.json", STYLES_DIR / raw, cand):
            if probe.is_file():
                cand = probe
                break
    if not cand.is_file():
        raise SystemExit(
            f"error: unknown --style {raw!r}. Built-in themes: "
            f"{', '.join(available_styles()) or '(none found)'}; "
            "or pass a path to a theme .json")
    try:
        data = json.loads(cand.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"error: cannot parse theme {cand}: {exc}")
    if not isinstance(data, dict):
        raise SystemExit(f"error: theme {cand} must be a JSON object")
    theme = Theme(data)
    if not data.get("id"):
        theme.id = cand.stem
    return theme


NUMERIC_RE = re.compile(
    r"^[\(\[-]?[$€£¥]?[\d,]+(?:\.[\d]+)?%?[\)\]]?$|"
    r"^(?:n\.?m\.?|NM|—|–|-|n/a)$", re.IGNORECASE)
CAPTION_RE = re.compile(r"^(?:Exhibit|Table|Chart)\s*\d*[:.\-]?\s*(.+)$", re.I)
SOURCE_RE  = re.compile(r"^\s*Source\s*:", re.I)


# ----------------------------------------------------------------------------
# Minimal YAML-frontmatter parser (flat keys, list-of-maps, folded scalars)
# ----------------------------------------------------------------------------
def parse_frontmatter(text: str) -> tuple[dict, str]:
    """Return (meta, body). meta = {} if no frontmatter block."""
    if not text.startswith("---"):
        return {}, text
    lines = text.splitlines(keepends=True)
    end = None
    for i in range(1, min(len(lines), 400)):
        if lines[i].strip() == "---":
            end = i
            break
    if end is None:
        return {}, text
    meta: dict = {}
    fm = lines[1:end]
    body = "".join(lines[end + 1:])
    i = 0
    list_key: str | None = None
    cur_item: dict | None = None
    while i < len(fm):
        line = fm[i].rstrip("\n")
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            i += 1
            continue
        if list_key is not None and stripped.startswith("- "):
            # new list item: `- key: value, key2: value2`
            cur_item = {}
            for pair in _split_pairs(stripped[2:]):
                k, v = pair
                cur_item[k] = _unquote(v)
            meta[list_key].append(cur_item)
            i += 1
            continue
        if list_key is not None and cur_item is not None and line[:1] in (" ", "\t"):
            # continuation line of the current list item (e.g. `title: ...`)
            m2 = re.match(r"^(\s*)([A-Za-z0-9_][\w]*)\s*:\s*(.*)$", line)
            if m2:
                cur_item[m2.group(2)] = _unquote(m2.group(3))
                i += 1
                continue
            i += 1
            continue
        m = re.match(r"^([A-Za-z0-9_][\w]*)\s*:\s*(.*)$", line)
        if not m:
            i += 1
            continue
        key, val = m.group(1).strip(), m.group(2).strip()
        if val == "":
            # valueless key: treat as an empty list (list items follow indented)
            meta[key] = []
            list_key = key
            i += 1
            continue
        if val.startswith(">-") or val.startswith("|-"):
            # folded / literal block: consume following indented lines
            chunks = []
            j = i + 1
            while j < len(fm) and (fm[j].startswith("  ") or fm[j].startswith("\t")):
                chunks.append(fm[j].strip())
                j += 1
            sep = " " if val.startswith(">-") else "\n"
            meta[key] = sep.join(chunks)
            list_key = None
            i = j
            continue
        if val == "[":
            meta[key] = []
            list_key = key
            i += 1
            continue
        meta[key] = _unquote(val)
        list_key = None
        i += 1
    return meta, body


def _split_pairs(s: str) -> list[tuple[str, str]]:
    """Split 'a: 1, b: 2' style inline pairs (used for list items)."""
    out = []
    for part in re.split(r",\s*(?=[A-Za-z_][\w]*\s*:)", s):
        if ":" in part:
            k, v = part.split(":", 1)
            out.append((k.strip(), v.strip()))
    return out


def _unquote(v: str):
    v = v.strip()
    if len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
        return v[1:-1]
    return v


# ----------------------------------------------------------------------------
# docx low-level helpers
# ----------------------------------------------------------------------------
def _set_font(run, name="Calibri", size=10, bold=False, italic=False,
              color=None, caps=False, eastasia=EASTASIA):
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    if color is not None:
        f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    # drop theme attributes: asciiTheme would override w:ascii in Word
    for t in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        attr = qn(t)
        if rFonts.get(attr) is not None:
            del rFonts.attrib[attr]
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), eastasia)
    if caps:
        c = OxmlElement("w:caps")
        c.set(qn("w:val"), "1")
        rPr.append(c)


def _set_char_spacing(run, twips: int):
    rPr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(twips))
    rPr.append(sp)


def _para_border(p, edge="bottom", color=RULEGRAY, sz=4, space=2):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def _shade_cell(cell, hexfill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _cell_border(cell, edges=("top", "bottom", "left", "right"), color="1F3864",
                 sz=6, edge_specs=None):
    """Apply cell borders. `edge_specs` = {edge: (hex_color, sz)} merged over `edges`."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    specs = {e: (color, sz) for e in edges}
    if edge_specs:
        specs.update(edge_specs)
    for e in ("top", "bottom", "left", "right", "insideH", "insideV"):
        if e not in specs:
            continue
        c, s = specs[e]
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(s))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), c)
        borders.append(el)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _add_field(paragraph, code: str):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    return run


def _page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


# ----------------------------------------------------------------------------
# Inline content: list of (text, fmt-dict)
# ----------------------------------------------------------------------------
class MDWalker(HTMLParser):
    """Convert markdown-lib HTML output into a list of block dicts."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[dict] = []
        self.stack: list[str] = []
        self.inlines: list = []          # current inline accumulator
        self.in_list = False
        self.list_type = None
        self.list_items: list[list] = []
        # table state
        self.in_table = False
        self.table_rows: list[tuple[bool, list[list], list]] = []
        self.table_aligns: list[str] = []
        self.cur_row = None
        self.cur_cell = None
        self.cur_cell_align = None
        self.in_thead = False
        # image handling
        self.cur_img = None

    # -- helpers ------------------------------------------------------------
    def _fmt(self) -> dict:
        s = self.stack
        f = {"bold": ("strong" in s) or ("b" in s),
             "italic": ("em" in s) or ("i" in s),
             "code": "code" in s}
        if "a" in s:
            f["link"] = True
        return f

    def _emit_para(self, kind="para", **kw):
        if self.inlines or kind in ("img",):
            self.blocks.append({"type": kind, "inlines": self.inlines, **kw})
        self.inlines = []

    # -- events -------------------------------------------------------------
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("h1", "h2", "h3", "h4", "h5"):
            self._emit_para()
            self.stack.append(tag)
            self.inlines = []
        elif tag == "p":
            if self.in_list and "li" in self.stack:
                # <p> inside a list item: keep accumulating into the item
                self.stack.append("p")
                return
            self._emit_para()
            self.stack.append("p")
            self.inlines = []
        elif tag == "blockquote":
            self._emit_para()
            self.stack.append("blockquote")
            self.inlines = []
        elif tag == "pre":
            self._emit_para()
            self.stack.append("pre")
            self.inlines = []
        elif tag == "ul" or tag == "ol":
            self._emit_para()
            self.in_list = True
            self.list_type = tag
            self.list_items = []
            self.stack.append(tag)
        elif tag == "li":
            self._emit_para()
            self.stack.append("li")
            self.inlines = []
        elif tag == "table":
            self._emit_para()
            self.in_table = True
            self.table_rows = []
            self.table_aligns = []
            self.in_thead = False
            self.stack.append("table")
        elif tag == "thead":
            self.in_thead = True
        elif tag == "tbody":
            self.in_thead = False
        elif tag == "tr":
            if self.in_table:
                self.cur_row = []
        elif tag in ("td", "th"):
            if self.in_table and self.cur_row is not None:
                self.cur_cell = []
                self.cur_cell_align = a.get("align")
                self.stack.append(tag)
                if tag == "th" and a.get("align"):
                    self.table_aligns.append(a["align"])
        elif tag == "img":
            src = a.get("src", "")
            alt = a.get("alt", "")
            # images arrive inside <p>; close current para content first
            self._emit_para(kind="img", src=src, alt=alt)
        elif tag == "hr":
            self._emit_para(kind="hr")
        elif tag == "br":
            self.inlines.append(("\n", self._fmt()))
        elif tag == "div":
            cls = a.get("class", "")
            if "pagebreak" in cls:
                self.blocks.append({"type": "pagebreak"})
        else:
            self.stack.append(tag)

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3", "h4", "h5"):
            lvl = int(tag[1])
            self._emit_para(kind="heading", level=lvl)
            if tag in self.stack:
                self.stack.remove(tag)
        elif tag == "p":
            if self.in_list and "li" in self.stack:
                if "p" in self.stack:
                    self.stack.remove("p")
                return
            self._emit_para()
            if "p" in self.stack:
                self.stack.remove("p")
        elif tag == "blockquote":
            self._emit_para(kind="quote")
            if "blockquote" in self.stack:
                self.stack.remove("blockquote")
        elif tag == "pre":
            self._emit_para(kind="codeblock")
            if "pre" in self.stack:
                self.stack.remove("pre")
        elif tag == "li":
            if self.inlines:
                self.list_items.append(self.inlines)
            self.inlines = []
            if "li" in self.stack:
                self.stack.remove("li")
        elif tag == "ul" or tag == "ol":
            self._emit_para()
            if self.list_items:
                self.blocks.append({"type": "list", "list_type": self.list_type,
                                    "items": self.list_items})
            self.in_list = False
            self.list_type = None
            self.list_items = []
            if tag in self.stack:
                self.stack.remove(tag)
        elif tag == "td" or tag == "th":
            if self.in_table and self.cur_cell is not None:
                self.cur_row.append((tag == "th", self.cur_cell, self.cur_cell_align))
                self.cur_cell = None
                self.cur_cell_align = None
            if tag in self.stack:
                self.stack.remove(tag)
        elif tag == "tr":
            if self.in_table and self.cur_row is not None:
                is_head = self.in_thead or all(c[0] for c in self.cur_row if c)
                self.table_rows.append((is_head,
                                        [c[1] for c in self.cur_row],
                                        [c[2] for c in self.cur_row]))
                self.cur_row = None
        elif tag == "table":
            self.in_table = False
            self.blocks.append({"type": "table", "rows": self.table_rows,
                                "aligns": self.table_aligns})
            if "table" in self.stack:
                self.stack.remove("table")
        elif tag == "thead":
            self.in_thead = False
        elif tag == "a":
            if "a" in self.stack:
                self.stack.remove("a")
        elif tag in self.stack:
            self.stack.remove(tag)

    def handle_data(self, data):
        if not data:
            return
        if self.in_table:
            # Only cell content counts; the whitespace/newlines BETWEEN table
            # tags (<table>, <thead>, <tr>, ...) must never leak into the
            # inline accumulator — otherwise every table produces a phantom
            # paragraph of 20+ line breaks that pushes the following content
            # onto the next page and can leave blank pages behind.
            if self.cur_cell is not None:
                self.cur_cell.append((data, self._fmt()))
            return
        if self.in_list and "li" in self.stack:
            self.inlines.append((data, self._fmt()))
            return
        # Only append text that actually belongs to a content context.
        # Whitespace between tags (e.g. "\n" between </li> and <li>, or after
        # a </table>) is structural, not content, and must be discarded.
        if "p" in self.stack or "li" in self.stack:
            self.inlines.append((data, self._fmt()))
            return
        if ("pre" in self.stack or "blockquote" in self.stack
                or any(h in self.stack for h in ("h1", "h2", "h3", "h4", "h5"))):
            self.inlines.append((data, self._fmt()))


# ----------------------------------------------------------------------------
# Renderer
# ----------------------------------------------------------------------------
class DocxRenderer:
    def __init__(self, meta: dict, md_path: Path, opts, theme: Theme | None = None):
        self.meta = meta
        self.md_dir = md_path.parent
        self.opts = opts
        self.th = theme or Theme({})
        self.exhibit_no = 0
        self.h1_no = 0
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    # -- theme helpers ------------------------------------------------------
    def _font(self, run, name=None, size=None, bold=False, italic=False,
              color=None, caps=False, mono=False, heading=False):
        """Theme-aware _set_font: defaults to the theme body/heading/mono face."""
        if name is None:
            name = self.th.mono_font if mono else (
                self.th.heading_font if heading else self.th.body_font)
        if size is None:
            size = self.th.body_size
        _set_font(run, name=name, size=size, bold=bold, italic=italic,
                  color=color, caps=caps)

    def _meta_case(self, text: str) -> str:
        return text.upper() if self.th.uppercase_meta else text

    # -- document scaffolding ----------------------------------------------
    def _setup_page(self):
        th = self.th
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(th.margin_lr)
        sec.top_margin, sec.bottom_margin = Inches(th.margin_tb), Inches(th.margin_tb)

    def _setup_styles(self):
        th = self.th
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = th.body_font
        normal.font.size = Pt(th.body_size)
        self._style_rfonts(normal.element.get_or_add_rPr(), th.body_font)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = th.line_spacing

        for name, size, bold, color, before, after, fname in (
                ("Heading 1", th.h1_size, True, th.navy, 16, 8, th.h1_font_name),
                ("Heading 2", th.h2_size, True, th.accent, 10, 4, th.heading_font),
                ("Heading 3", th.h3_size, True, th.text, 8, 3, th.heading_font),
                ("Heading 4", th.h4_size, True, th.midgray, 6, 2, th.heading_font)):
            h = st[name]
            h.font.name = fname; h.font.size = Pt(size); h.font.bold = bold
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            self._style_rfonts(h.element.get_or_add_rPr(), fname)

    @staticmethod
    def _style_rfonts(rpr, name, eastasia=EASTASIA):
        """Ensure EXACTLY ONE w:rFonts element with ascii/hAnsi/eastAsia set.

        python-docx's font.name setter creates <w:rFonts w:ascii w:hAnsi> but
        appends nothing for eastAsia; appending a second rFonts leaves the
        East-Asian face unregistered (Word picks whichever duplicate it reads
        first). Instead: reuse the existing element and set all three faces.
        """
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for t in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            attr = qn(t)
            if rf.get(attr) is not None:
                del rf.attrib[attr]
        rf.set(qn("w:ascii"), name)
        rf.set(qn("w:hAnsi"), name)
        rf.set(qn("w:eastAsia"), eastasia)

    # -- cover --------------------------------------------------------------
    def render_cover(self):
        m = self.meta
        th = self.th
        firm = m.get("firm", "Independent Equity Research")
        rtype = m.get("report_type", "Equity Research")
        date = m.get("date", "")

        # top line: firm | date
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(self._meta_case(firm))
        self._font(r, size=9, bold=True, color=th.accent, caps=th.uppercase_meta)
        _set_char_spacing(r, 14)
        if date:
            r2 = p.add_run(f"\t{date}")
            self._font(r2, size=9, color=th.midgray)
        if th.masthead_rule:
            _para_border(p, "bottom", th.rule, sz=6)

        # logo
        if m.get("logo"):
            logo = self.md_dir / m["logo"]
            if logo.exists():
                self.doc.add_picture(str(logo), width=Inches(1.6))
                self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT

        # report-type banner
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(self._meta_case("Equity Research"))
        self._font(r, size=8, color=th.midgray, caps=th.uppercase_meta)
        _set_char_spacing(r, 18)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(self._meta_case(rtype))
        self._font(r, size=15, bold=True, color=th.accent, caps=th.uppercase_meta)
        _set_char_spacing(r, 16)
        if th.masthead_rule:
            _para_border(p, "bottom", th.accent_hex, sz=10)

        # company
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(m.get("company", ""))
        self._font(r, size=25, bold=True, color=th.navy)
        ident = m.get("ticker", "")
        if m.get("exchange"):
            ident += f" · {m['exchange']}"
        if m.get("sector"):
            ident += f" · {m['sector']}"
        p2 = self.doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        r = p2.add_run(ident)
        self._font(r, size=10.5, color=th.midgray)

        # title
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(m.get("title", ""))
        self._font(r, size=13.5, bold=True, color=th.text, italic=True)

        self._render_rating_box()
        self._render_key_data_line()
        if m.get("price_chart"):
            pc = self.md_dir / m["price_chart"]
            if pc.exists():
                self.exhibit_no += 1
                self._render_figure(pc, m.get("alt_price_chart",
                                              "12-month price performance"))

        self._render_analysts()
        self._render_conflict_line()
        # footer line of cover
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run(f"© {m.get('copyright_year', '')} {firm} · "
                      "See important disclosures on the final pages of this report.")
        self._font(r, size=7.5, color=th.lightgray)

    def _render_rating_box(self):
        m = self.meta
        th = self.th
        labels = ["Rating", "Target Price", "Current Price", "Upside/(Downside)"]
        values = [m.get("rating", "—"),
                  m.get("target_price_prev") and
                  f"{m['currency']} {m.get('target_price_prev')} → {m['currency']} {m.get('target_price')}"
                  or f"{m.get('currency', '$')} {m.get('target_price', '—')}",
                  f"{m.get('currency', '$')} {m.get('current_price', '—')}",
                  m.get("upside", "—")]
        rating_color = th.navy
        for k, c in th.rating_colors.items():
            if k in m.get("rating", "").lower():
                rating_color = c
                break
        style = th.rating_style
        table = self.doc.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for j, (lab, val) in enumerate(zip(labels, values)):
            for i in (0, 1):
                cell = table.cell(i, j)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _cell_margins(cell)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if style == "open":
                    if i == 0:
                        r = p.add_run(self._meta_case(lab))
                        self._font(r, size=8, bold=True, color=th.midgray,
                                   caps=th.uppercase_meta)
                        _set_char_spacing(r, 8)
                        _cell_border(cell, ("bottom",), th.rule, sz=6)
                    else:
                        r = p.add_run(val)
                        self._font(r, size=12, bold=True,
                                   color=rating_color if j == 0 else th.navy)
                        if th.rating_border:
                            _cell_border(cell, color=th.rule, sz=4)
                elif style == "rail":
                    _shade_cell(cell, th.rating_fill)
                    specs = {}
                    if j == 0:
                        specs["left"] = (th.accent_hex, 18)
                    if th.rating_border:
                        specs.update({"top": (th.rule, 4), "bottom": (th.rule, 4),
                                      "right": (th.rule, 4)})
                    if specs:
                        _cell_border(cell, edges=(), edge_specs=specs)
                    if i == 0:
                        r = p.add_run(self._meta_case(lab))
                        self._font(r, size=8, bold=True, color=th.midgray,
                                   caps=th.uppercase_meta)
                        _set_char_spacing(r, 8)
                    else:
                        r = p.add_run(val)
                        self._font(r, size=12, bold=True,
                                   color=rating_color if j == 0 else th.navy)
                else:  # "band": filled label band + value row (legacy default)
                    if i == 0:
                        _shade_cell(cell, th.table_header_fill)
                        r = p.add_run(self._meta_case(lab))
                        self._font(r, size=8, bold=True, color=th.table_header_text,
                                   caps=th.uppercase_meta)
                        _set_char_spacing(r, 8)
                    else:
                        _shade_cell(cell, th.rating_fill)
                        r = p.add_run(val)
                        self._font(r, size=12, bold=True,
                                   color=rating_color if j == 0 else th.navy)
                    if th.rating_border:
                        _cell_border(cell, color=th.table_header_fill, sz=6)
        if m.get("rating_note"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(m["rating_note"])
            self._font(r, size=8, italic=True, color=th.midgray)

    def _render_key_data_line(self):
        m = self.meta
        th = self.th
        bits = []
        if m.get("market_cap"):
            bits.append(f"Market Cap  {m['market_cap']}")
        if m.get("shares_out"):
            bits.append(f"Shares Out  {m['shares_out']}")
        if m.get("avg_daily_volume"):
            bits.append(f"Avg Daily Vol  {m['avg_daily_volume']}")
        if m.get("52wk_range"):
            bits.append(f"52-Week Range  {m['52wk_range']}")
        if m.get("pe_ntm"):
            bits.append(f"NTM P/E  {m['pe_ntm']}")
        if m.get("horizon"):
            bits.append(f"Horizon  {m['horizon']}")
        if not bits:
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(12)
        for i, b in enumerate(bits):
            if i:
                p.add_run("   ·   ")
            r = p.add_run(b)
            self._font(r, size=8.5, color=th.midgray)

    def _render_analysts(self):
        m = self.meta
        th = self.th
        analysts = m.get("analysts") or []
        if not analysts:
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(self._meta_case("Analysts"))
        self._font(r, size=8, bold=True, color=th.accent, caps=th.uppercase_meta)
        _set_char_spacing(r, 12)
        for a in analysts:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(a.get("name", ""))
            self._font(r, size=10.5, bold=True, color=th.navy)
            extra = " · ".join(x for x in (a.get("title"), a.get("phone"),
                                           a.get("email")) if x)
            if extra:
                r2 = p.add_run(f"   {extra}")
                self._font(r2, size=8.5, color=th.midgray)

    def _render_conflict_line(self):
        m = self.meta
        text = m.get("disclosure_conflict", "")
        if not text:
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        self._font(r, size=8, italic=True, color=self.th.midgray)

    # -- body ---------------------------------------------------------------
    def render_body(self, blocks):
        # post-process: caption / source classification
        for i, b in enumerate(blocks):
            if b["type"] == "para":
                txt = "".join(t for t, _ in b["inlines"]).strip()
                if SOURCE_RE.match(txt):
                    b["type"] = "source"
                elif (i + 1 < len(blocks) and blocks[i + 1]["type"] == "table"
                      and CAPTION_RE.match(txt)):
                    b["type"] = "caption"
        for b in blocks:
            t = b["type"]
            if t == "pagebreak":
                _page_break(self.doc)
            elif t == "heading":
                self._render_heading(b["level"], b["inlines"])
            elif t == "para":
                self._render_para(b["inlines"], indent=0)
            elif t == "quote":
                self._render_quote(b["inlines"])
            elif t == "source":
                self._render_para(b["inlines"], size=7.5, italic=True,
                                  color=self.th.midgray, space_after=8)
            elif t == "caption":
                self._render_para(b["inlines"], size=8.5, bold=True,
                                  color=self.th.midgray, space_after=2,
                                  keep_with_next=True)
            elif t == "list":
                self._render_list(b["list_type"], b["items"])
            elif t == "table":
                self._render_table(b["rows"], b["aligns"])
            elif t == "img":
                self._render_img(b["src"], b["alt"])
            elif t == "codeblock":
                self._render_codeblock(b["inlines"])
            elif t == "hr":
                p = self.doc.add_paragraph()
                _para_border(p, "bottom", self.th.rule, sz=4)

    def _render_heading(self, level, inlines):
        th = self.th
        text = self._sub("".join(t for t, _ in inlines))
        if level == 1 and not self.opts.no_numbering:
            self.h1_no += 1
            text = f"{self.h1_no}.  {text}"
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        r = p.add_run(text)
        if level == 1:
            self._font(r, name=th.h1_font_name, size=th.h1_size, bold=True,
                       color=th.navy)
            _para_border(p, "bottom", th.rule, sz=4)
        elif level == 2:
            self._font(r, heading=True, size=th.h2_size, bold=True, color=th.accent)
        elif level == 3:
            self._font(r, heading=True, size=th.h3_size, bold=True, color=th.text)
        else:
            self._font(r, heading=True, size=th.h4_size, bold=True,
                       color=th.midgray, italic=True)

    def _render_para(self, inlines, size=None, bold=False, italic=False,
                     color=None, indent=0, space_after=6, keep_with_next=False):
        if size is None:
            size = self.th.body_size
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        if indent:
            pf.left_indent = Inches(indent)
        if keep_with_next:
            pf.keep_with_next = True
        self._write_inlines(p, inlines, size=size, base_bold=bold,
                            base_italic=italic, base_color=color)
        return p

    def _write_inlines(self, p, inlines, size=None, base_bold=False,
                       base_italic=False, base_color=None):
        if size is None:
            size = self.th.body_size
        for text, f in inlines:
            text = self._sub(text)
            if "\n" in text:
                parts = text.split("\n")
                for k, part in enumerate(parts):
                    if k:
                        p.add_run().add_break()
                    if part:
                        self._add_run(p, part, size=size, bold=base_bold or f.get("bold"),
                                      italic=base_italic or f.get("italic"),
                                      color=base_color, code=f.get("code"),
                                      link=f.get("link"))
            else:
                self._add_run(p, text, size=size, bold=base_bold or f.get("bold"),
                              italic=base_italic or f.get("italic"),
                              color=base_color, code=f.get("code"), link=f.get("link"))

    def _add_run(self, p, text, size=None, bold=False, italic=False, color=None,
                 code=False, link=False):
        th = self.th
        if size is None:
            size = th.body_size
        r = p.add_run(text)
        if code:
            self._font(r, name=th.mono_font, size=max(size - 1, 7),
                       color=th.code_text)
        elif link:
            self._font(r, size=size, color=th.link)
            r.font.underline = True
        else:
            self._font(r, size=size, bold=bold, italic=italic, color=color)
        return r

    def _render_quote(self, inlines):
        th = self.th
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3)
        pf.space_after = Pt(6)
        _para_border(p, "left", th.accent_hex, sz=12, space=6)
        self._write_inlines(p, inlines, size=9.5, base_italic=True,
                            base_color=th.midgray)

    def _render_list(self, list_type, items):
        style = "List Number" if list_type == "ol" else "List Bullet"
        for item in items:
            p = self.doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(3)
            self._write_inlines(p, item, size=self.th.body_size)

    def _sub(self, text: str) -> str:
        """Replace {{FIELD}} tokens with frontmatter values (case-insensitive)."""
        if "{{" not in text:
            return text

        def rep(m):
            key = m.group(1).strip().lower()
            for k, v in self.meta.items():
                if k.lower() == key:
                    if isinstance(v, list):
                        return str(v[0].get("name", "")) if v else ""
                    return str(v)
            return m.group(0)

        return re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", rep, text)

    def _is_numeric_col(self, rows, j):
        vals = []
        for is_head, cells, _ in rows:
            if is_head or j >= len(cells):
                continue
            txt = "".join(t for t, _ in cells[j]).strip()
            vals.append(txt)
        if not vals:
            return False
        hits = sum(1 for v in vals if NUMERIC_RE.match(v))
        return hits / len(vals) >= 0.6

    def _render_table(self, rows, aligns):
        if not rows:
            return
        th = self.th
        # keep a short lead-in paragraph (e.g. "**Catalyst calendar:**") on the same
        # page as the table's first row — prevents orphaned lead-ins at page bottom
        if self.doc.paragraphs:
            last_p = self.doc.paragraphs[-1]
            if 0 < len(last_p.text.strip()) <= 150:
                last_p.paragraph_format.keep_with_next = True
        ncols = max(len(cells) for _, cells, _ in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        col_w = Inches(6.6 / ncols)
        numeric_cols = [self._is_numeric_col(rows, j) for j in range(ncols)]
        has_head = any(is_head for is_head, _, _ in rows)
        # repeat the header row on every page a table spans (pagination sanity)
        trPr = table.rows[0]._tr.get_or_add_trPr()
        th_el = OxmlElement("w:tblHeader")
        th_el.set(qn("w:val"), "true")
        trPr.append(th_el)
        for i, (is_head, cells, cell_aligns) in enumerate(rows):
            for j in range(ncols):
                cell = table.cell(i, j)
                cell.width = col_w
                _cell_margins(cell, top=30, bottom=30, left=70, right=70)
                if is_head:
                    _shade_cell(cell, th.table_header_fill)
                elif th.zebra and th.table_zebra and i % 2 == 0:
                    _shade_cell(cell, th.table_zebra)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if numeric_cols[j]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif j < len(cell_aligns) and cell_aligns[j] == "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif j < len(cell_aligns) and cell_aligns[j] == "right":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                inlines = cells[j] if j < len(cells) else []
                self._write_inlines(p, inlines, size=th.table_size,
                                    base_bold=is_head,
                                    base_color=th.table_header_text if is_head else None)
        # grid: hairline horizontal rules always, verticals only when themed on
        wanted = {"top", "bottom", "insideH"}
        if th.vertical_rules:
            wanted |= {"left", "right", "insideV"}
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}")
            if e in wanted:
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), th.table_grid)
            else:
                el.set(qn("w:val"), "none")
                el.set(qn("w:sz"), "0")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "auto")
            borders.append(el)
        tblPr.append(borders)
        # theme header rules: stronger line above and/or below the header row
        if has_head and (th.rule_above_header or th.rule_below_header):
            edges = []
            if th.rule_above_header:
                edges.append("top")
            if th.rule_below_header:
                edges.append("bottom")
            for j in range(ncols):
                _cell_border(table.cell(0, j), tuple(edges), th.navy_hex, sz=12)

    def _render_img(self, src, alt):
        img_path = self.md_dir / src if src else None
        if not img_path or not img_path.exists():
            p = self.doc.add_paragraph()
            r = p.add_run(f"[missing image: {src}]")
            self._font(r, size=9, italic=True, color=self.th.lightgray)
            return
        self.exhibit_no += 1
        self._render_figure(img_path, alt)

    def _render_figure(self, img_path, alt):
        # caption above (WS convention: "Exhibit N: <title>")
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        cap = alt or "Price performance"
        r = p.add_run(f"Exhibit {self.exhibit_no}:  {cap}")
        self._font(r, size=8.5, bold=True, color=self.th.midgray)
        pic_p = self.doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_p.paragraph_format.space_after = Pt(2)
        run = pic_p.add_run()
        run.add_picture(str(img_path), width=Inches(5.8))

    def _render_codeblock(self, inlines):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(6)
        # shading
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), self.th.code_fill)
        pPr.append(shd)
        self._write_inlines(p, inlines,
                            size=max(round(self.th.body_size - 2, 1), 7.0),
                            base_color=None)

    # -- header / footer (body section) -------------------------------------
    def add_body_section(self):
        m = self.meta
        th = self.th
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(th.margin_lr)
        sec.top_margin, sec.bottom_margin = Inches(th.margin_tb), Inches(th.margin_tb)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False

        # header: title left, date right, thin rule
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_ALIGN_PARAGRAPH.RIGHT)
        r = hp.add_run(m.get("header_title") or f"{m.get('company', '')} ({m.get('ticker', '')})")
        self._font(r, size=8, bold=True, color=th.midgray, caps=th.uppercase_meta)
        r = hp.add_run(f"\t{m.get('date', '')}")
        self._font(r, size=8, color=th.lightgray)
        if th.masthead_rule:
            _para_border(hp, "bottom", th.rule, sz=4)

        # footer: org left, "Page X of Y" right
        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_ALIGN_PARAGRAPH.RIGHT)
        r = fp.add_run(m.get("footer_org", ""))
        self._font(r, size=7.5, color=th.lightgray)
        r = fp.add_run("\tPage ")
        self._font(r, size=7.5, color=th.lightgray)
        _add_field(fp, "PAGE")
        r = fp.add_run(" of ")
        self._font(r, size=7.5, color=th.lightgray)
        _add_field(fp, "NUMPAGES")
        return sec

    def add_toc(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run("Table of Contents")
        self._font(r, size=13, bold=True, color=self.th.navy)
        p2 = self.doc.add_paragraph()
        run = p2.add_run()
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\o "1-2" \\h \\z \\u '
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = "Right-click and choose 'Update Field' to build the table of contents."
        fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
        run._r.append(t); run._r.append(fld3)
        _page_break(self.doc)


# ----------------------------------------------------------------------------
def convert(md_path: Path, out_path: Path, opts, theme: Theme | None = None) -> Path:
    text = md_path.read_text(encoding="utf-8")
    meta, body = parse_frontmatter(text)
    html = md_lib.markdown(body, extensions=["tables", "fenced_code", "sane_lists"])
    walker = MDWalker()
    walker.feed(html)
    walker.close()

    r = DocxRenderer(meta, md_path, opts, theme=theme)
    r.render_cover()
    r.add_body_section()
    if opts.toc:
        r.add_toc()
    r.render_body(walker.blocks)
    r.doc.save(str(out_path))
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description="Wall Street-style Markdown → docx")
    ap.add_argument("input", help="input .md report (templates/report_template.md)")
    ap.add_argument("-o", "--output", help="output .docx path")
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF via docx2pdf (needs MS Word)")
    ap.add_argument("--toc", action="store_true", help="insert a Table of Contents field")
    ap.add_argument("--no-numbering", action="store_true", help="do not number H1 sections")
    ap.add_argument("--style", metavar="ID|PATH",
                    help="institution style theme: an id under templates/styles/ "
                         "(goldman_hardline, morganstanley_restrained, jpmorgan_heavyset, "
                         "barclays_cyanline, bernstein_monochrome, ubs_swissminimal) or a "
                         "path to a theme .json; default = built-in Wall Street look")
    opts = ap.parse_args(argv)

    src = Path(opts.input)
    if not src.exists():
        sys.exit(f"error: input not found: {src}")
    out = Path(opts.output) if opts.output else src.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    theme = load_theme(opts.style)
    convert(src, out, opts, theme=theme)
    print(f"[ok] docx written: {out} ({out.stat().st_size:,} bytes)")
    if opts.style:
        print(f"[ok] style: {theme.id} ({theme.display_name})")

    if opts.pdf:
        pdf = out.with_suffix(".pdf")
        # Word/COM conversion is flaky when invoked repeatedly in quick
        # succession (a stale Word instance can leave the previous file on
        # disk while the call returns cleanly). Retry a couple of times and
        # verify freshness instead of trusting the call.
        import time as _time
        last_err = None
        for attempt in (1, 2, 3):
            try:
                from docx2pdf import convert as d2p
                d2p(str(out), str(pdf))
            except Exception as e:  # noqa: BLE001
                last_err = e
            if pdf.exists() and pdf.stat().st_mtime >= out.stat().st_mtime:
                print(f"[ok] pdf written: {pdf} ({pdf.stat().st_size:,} bytes)")
                return 0
            if attempt < 3:
                _time.sleep(2.0)
        print(f"[warn] PDF export failed after 3 attempts (last error: {last_err}).\n"
              "  The .docx is current but the PDF is stale or missing. Re-run, or export "
              "manually from Word.\n"
              "  Alternatives: (1) open the .docx in Microsoft Word and Save As PDF, "
              "(2) install pandoc + a LaTeX engine, or (3) install typst.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
